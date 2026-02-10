import io
import logging
from fastapi import APIRouter, Request, HTTPException, Form
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from typing import List, Optional

from ..models.schemas import LicenseRequest, LicenseResponse
from ...core.license_generator import generate_license_text

logger = logging.getLogger(__name__)

router = APIRouter()

# Setup Jinja2 templates
templates = Jinja2Templates(directory="src/templates")


@router.get("/", response_class=HTMLResponse)
async def index(request: Request):
    """Render the main license generation form"""
    return templates.TemplateResponse(request, "index.html")


@router.post("/generate_license", response_class=HTMLResponse)
async def generate_license(
    request: Request,
    organization_name: str = Form(...),
    copyright_holder: str = Form(...),
    scope: List[str] = Form(...),
    boundary: str = Form(...),
    territory: List[str] = Form(default_factory=list),
    other_territory: Optional[str] = Form(None),
    attribution: str = Form(...),
    distribution: str = Form(...),
    llm: str = Form(...),
    warranty: str = Form(...),
    authbody: str = Form(...),
    Authbody_name: Optional[str] = Form(None),
):
    """
    Generate license from form data.
    Accepts form-encoded data and returns HTML response (for HTMX compatibility).
    """
    try:
        license_req = LicenseRequest(
            organization_name=organization_name,
            copyright_holder=copyright_holder,
            scope=scope,
            boundary=boundary,
            territory=territory,
            other_territory=other_territory,
            attribution=attribution,
            distribution=distribution,
            llm=llm,
            warranty=warranty,
            authbody=authbody,
            Authbody_name=Authbody_name,
        )

        logger.info('Generating license for organization: %s', license_req.organization_name)

        form_data = license_req.model_dump()
        license_text, human_readable = generate_license_text(form_data)

        # Store both in session so download routes don't need to regenerate
        request.session['license_text'] = license_text
        request.session['form_data'] = form_data

        context = {
            "clauses_text": human_readable,
            "license_text": license_text,
        }

        # HTMX requests get a partial template; regular form submits get the full page
        is_htmx = request.headers.get("HX-Request") == "true"
        template = "components/license_result.html" if is_htmx else "result.html"

        return templates.TemplateResponse(request, template, context)

    except ValueError as e:
        logger.error('Validation error: %s', str(e))
        return templates.TemplateResponse(
            request, "error.html",
            {"message": f"Invalid form data: {str(e)}"},
            status_code=400
        )
    except Exception as e:
        logger.error('Error generating license: %s', str(e))
        return templates.TemplateResponse(
            request, "error.html",
            {"message": f"Error generating license: {str(e)}"},
            status_code=500
        )


def _get_license_text_from_session(request: Request) -> str:
    """Retrieve license text from session, raising 400 if not found."""
    license_text = request.session.get('license_text')
    if not license_text:
        raise HTTPException(status_code=400, detail="No license data found. Please generate a license first.")
    return license_text


@router.get("/download/md")
async def download_md(request: Request):
    """Download the generated license as a Markdown file"""
    license_text = _get_license_text_from_session(request)

    file_obj = io.BytesIO(license_text.encode('utf-8'))

    return StreamingResponse(
        file_obj,
        media_type="text/markdown",
        headers={"Content-Disposition": "attachment; filename=License.md"}
    )


@router.get("/download/pdf")
async def download_pdf(request: Request):
    """Download the generated license as a PDF file"""
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    license_text = _get_license_text_from_session(request)

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_left_margin(20)
    pdf.set_right_margin(20)
    effective_width = pdf.w - pdf.l_margin - pdf.r_margin

    pdf.set_font("Helvetica", style="B", size=14)
    pdf.cell(effective_width, 10, "SOFTWARE LICENSE AGREEMENT",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
    pdf.ln(4)
    pdf.set_font("Helvetica", size=9)

    # Skip first line if it's the title already rendered above
    body_lines = license_text.split('\n')
    start = 1 if body_lines and body_lines[0].strip() == "SOFTWARE LICENSE AGREEMENT" else 0
    body_text = '\n'.join(body_lines[start:])

    pdf.multi_cell(effective_width, 5, body_text)

    pdf_bytes = pdf.output()

    return StreamingResponse(
        io.BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=License.pdf"}
    )


@router.get("/download/docx")
async def download_docx(request: Request):
    """Download the generated license as a Word (.docx) file"""
    from docx import Document
    from docx.shared import Pt

    license_text = _get_license_text_from_session(request)

    doc = Document()

    # Style the title
    title = doc.add_heading("SOFTWARE LICENSE AGREEMENT", level=0)
    title.alignment = 1  # center

    # Add body — skip first line if it's the repeated title
    lines = license_text.split('\n')
    start = 1 if lines and lines[0].strip() == "SOFTWARE LICENSE AGREEMENT" else 0
    for line in lines[start:]:
        para = doc.add_paragraph(line)
        if para.runs:
            para.runs[0].font.size = Pt(10)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=License.docx"}
    )


# JSON API endpoint for programmatic access
@router.post("/api/generate", response_model=LicenseResponse)
async def api_generate_license(license_req: LicenseRequest):
    """
    API endpoint for license generation (returns JSON).
    Use this for programmatic access to the license generator.
    """
    try:
        logger.info('API: Generating license for organization: %s', license_req.organization_name)

        form_data = license_req.model_dump()
        license_text, human_readable = generate_license_text(form_data)

        return LicenseResponse(
            license_text=license_text,
            human_readable=human_readable,
        )

    except Exception as e:
        logger.error('API: Error generating license: %s', str(e))
        raise HTTPException(status_code=500, detail=f"Error generating license: {str(e)}")
